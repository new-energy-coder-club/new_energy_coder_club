# -*- coding: utf-8 -*-
"""在立项 PPT 第 4 页后插入'知识产权布局'页（风格与 slide4 一致）"""
import os, copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SRC = '/mnt/d/OneDrive/Desktop/1.工商业管理项目立项-解决方案AI项目(含SolarGlyph平台介绍).pptx'
OUT = '/mnt/d/Downloads/1.工商业管理项目立项-解决方案AI项目(含SolarGlyph平台介绍)_含专利布局.pptx'
FIG6 = '/mnt/d/Project_env/SolarGlyph/patent/figures/fig6_effect_compare.png'
FONT = '微软雅黑'

prs = Presentation(SRC)
layout = prs.slides[3].slide_layout  # 与第4页相同版式
slide = prs.slides.add_slide(layout)

def add_text(x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    return tb

def para(tf, text, size, bold=False, first=False, space_after=4):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(0, 0, 0)
    p.space_after = Pt(space_after)
    return p

# 页眉标签（沿用全 deck 风格）
tag = add_text(0.97, 0.36, 5.2, 0.71)
para(tag.text_frame, '项目知识产权布局', 36, bold=True, first=True)

# 正文
body = add_text(0.79, 0.88, 7.1, 6.3)
tf = body.text_frame
para(tf, '发明专利布局——基于图案掩膜的光伏阵列自动排布', 18, bold=True, first=True, space_after=8)
para(tf, '专利名称：一种基于图案掩膜的光伏阵列自动排布方法及系统。已完成技术交底书、权利要求书（14项）与查新检索，拟申报发明专利。', 13, space_after=8)

para(tf, '1. 核心创新点', 15, bold=True)
para(tf, '· 掩膜栅格分辨率与组件排布模数一一对应（“一格一组件”），图案像素即施工坐标；', 13)
para(tf, '· “掩膜定候选、工程约束定结果”两级自动仲裁：9检测点＋边界内缩＋障碍物缓冲逐块校验，无需人工修正；', 13)
para(tf, '· 冬至正午全向保守缓冲替代逐时阴影外包，单次计算即可在浏览器端实时完成排布。', 13, space_after=8)

para(tf, '2. 查新结论', 15, bold=True)
para(tf, '经检索国内外专利与文献，“图案掩膜自动排布”未见在先公开，具备新颖性与创造性空间；已识别并规避自动排布、运维通道、阴影外包类现有技术区。', 13, space_after=8)

para(tf, '3. 业务价值', 15, bold=True)
para(tf, '· 装机较小的客户可用屋顶阵列摆出企业标识，无人机航拍即可呈现，增强方案营销亮点；', 13)
para(tf, '· 消纳受限、不宜满铺的项目“不满铺也有亮点”，兼顾投资收益与品牌展示；', 13)
para(tf, '· 对比焱图/视觉动力等工具：无需无人机建模排队、图纸与项目数据不出企业。', 13, space_after=8)

para(tf, '4. 申报计划', 15, bold=True)
para(tf, '2026年Q3完成发明专利申报；同步储备“消纳反推测算”方向第二件专利，形成排布+测算的专利组合。', 13)

# 右侧配图 + 图注
slide.shapes.add_picture(FIG6, Inches(7.95), Inches(1.5), Inches(4.9))
cap = add_text(7.95, 4.15, 4.9, 0.7)
para(cap.text_frame, '掩膜排布（左，阵列呈现标识图案）与现有满铺排布（右）效果对比', 14, first=True)

# 将新页移动到第 4 页之后（索引 4）
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
new_id = ids[-1]
sldIdLst.remove(new_id)
sldIdLst.insert(4, new_id)

prs.save(OUT)
print('saved, total slides:', len(prs.slides._sldIdLst))
