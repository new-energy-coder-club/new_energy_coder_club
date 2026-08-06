# -*- coding: utf-8 -*-
"""红线资信信息列写入工具

读取 Excel 中指定主体列（默认"项目公司"），对每个去重主体调用
corporate-deep-query 脚本查询 5 个红线维度（失信/被执行/终本/行政处罚/经营异常），
将结果写入表格最右侧新增的"红线资信信息"列（样式复制自原最后一列）。

用法：
  python redline_credit_column.py --file 路径.xlsx \
      [--sheet 工作表名] [--entity-col 项目公司] \
      [--codes '{"公司名":"统一社会信用代码", ...}'] \
      [--col-name 红线资信信息] [--env prod]

说明：
  - --codes 可选。已知名称->代码映射时传入可跳过 qcc 查代码步骤；
    未提供时仅输出待查主体清单（--list-only 效果），由调用方（Agent）用
    qcc MCP get_company_by_query 逐个锁定代码后重跑。
  - 查询脚本路径默认取 corporate-deep-query skill 内置脚本，可用 --query-script 覆盖。
"""
import argparse
import json
import os
import re
import subprocess
import sys
from copy import copy

DEFAULT_QUERY_SCRIPT = os.path.expanduser(
    r"~\.config\agents\skills\corporate-deep-query\scripts\query_corporate.py")

MODULES = [
    ('BreakFaithExecutorInfoList', '失信被执行人'),
    ('JudicialExecutorInfoList', '被执行人'),
    ('FinalCaseInfoList', '终本案件'),
    ('AdmPenalInfoList', '行政处罚'),
    ('AbnormalInfoList', '经营异常'),
]


def query_count(query_script, code, env, module):
    """返回 (条数:int|None, 摘要行:str)。None 表示查询失败。"""
    r = subprocess.run(
        [sys.executable, query_script, '--code', code, '--env', env, '--module', module],
        capture_output=True, text=True, timeout=120, encoding='utf-8', errors='replace')
    out = (r.stdout or '') + (r.stderr or '')
    m = re.search(r'===\s*.*?（(\d+)\s*条）', out)
    if m:
        return int(m.group(1)), out.strip()
    return None, out.strip()


def build_text(name, code, results, date_str):
    """results: [(label, count|None), ...]"""
    parts = []
    for label, cnt in results:
        if cnt is None:
            parts.append(f'{label}查询失败')
        elif cnt == 0:
            parts.append(f'无{label}记录' if not label.endswith('案件') else f'无{label}')
        else:
            parts.append(f'⚠ {label} {cnt} 条')
    return f'{name}（{code}）\n经查询（{date_str}）：' + '、'.join(parts)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--file', required=True)
    ap.add_argument('--sheet', default=None, help='默认第一个工作表')
    ap.add_argument('--entity-col', default='项目公司', help='主体列表头名')
    ap.add_argument('--codes', default=None, help='JSON: {"公司名":"信用代码"}')
    ap.add_argument('--col-name', default='红线资信信息')
    ap.add_argument('--env', default='prod')
    ap.add_argument('--query-script', default=DEFAULT_QUERY_SCRIPT)
    ap.add_argument('--date', default=None, help='查询日期，默认今天 YYYY-MM-DD')
    args = ap.parse_args()

    import openpyxl
    from datetime import date
    date_str = args.date or date.today().isoformat()

    wb = openpyxl.load_workbook(args.file)
    ws = wb[args.sheet] if args.sheet else wb.worksheets[0]

    # 定位主体列（第1行表头）
    entity_col = None
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(1, c).value or '').strip() == args.entity_col:
            entity_col = c
            break
    if entity_col is None:
        print(f'ERROR: 未找到表头 "{args.entity_col}"，现有表头: '
              f'{[ws.cell(1, c).value for c in range(1, ws.max_column + 1)]}')
        sys.exit(1)

    # 收集数据行（跳过含公式的汇总行）及去重主体
    rows, names = [], []
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, entity_col).value
        if v is None or str(v).strip().startswith('='):
            continue
        name = str(v).strip()
        rows.append((r, name))
        if name not in names:
            names.append(name)

    codes = json.loads(args.codes) if args.codes else {}
    missing = [n for n in names if n not in codes]
    if missing:
        print('NEED_CODES: 以下主体缺少统一社会信用代码，请先用 qcc get_company_by_query 锁定：')
        for n in missing:
            print(' -', n)
        sys.exit(2)

    # 逐主体查询 5 个红线维度
    text_by_name = {}
    for name in names:
        code = codes[name]
        results = []
        for mod, label in MODULES:
            cnt, _ = query_count(args.query_script, code, args.env, mod)
            results.append((label, cnt))
            print(f'[{name}] {label}: {"失败" if cnt is None else str(cnt) + " 条"}')
        text_by_name[name] = build_text(name, code, results, date_str)

    # 若目标列已存在则覆盖，否则追加到最右
    target = None
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(1, c).value or '').strip() == args.col_name:
            target = c
            break
    if target is None:
        target = ws.max_column + 1
    ref = target - 1 if target > 1 else 1  # 样式参考列

    from openpyxl.utils import get_column_letter
    h = ws.cell(1, target, value=args.col_name)
    hr = ws.cell(1, ref)
    h.font, h.fill, h.border, h.alignment = (copy(hr.font), copy(hr.fill),
                                             copy(hr.border), copy(hr.alignment))
    for r, name in rows:
        c = ws.cell(r, target, value=text_by_name[name])
        cr = ws.cell(r, ref)
        c.font, c.fill, c.border, c.alignment = (copy(cr.font), copy(cr.fill),
                                                 copy(cr.border), copy(cr.alignment))
    col_letter = get_column_letter(target)
    ref_letter = get_column_letter(ref)
    ws.column_dimensions[col_letter].width = ws.column_dimensions[ref_letter].width or 40

    wb.save(args.file)
    print(f'DONE: "{args.col_name}" 写入 {args.file} 第 {target} 列，共 {len(rows)} 行')


if __name__ == '__main__':
    main()
